from __future__ import annotations

from pathlib import Path
import os
from typing import Optional

import requests
from bs4 import BeautifulSoup
from pypdf import PdfReader

from app.core.config import settings
from app.services.url_guard import validate_web_url


def load_document(path: Path) -> list[tuple[Optional[int], str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        pages = [(index, page.extract_text() or "") for index, page in enumerate(reader.pages, 1)]
        if any(text.strip() for _, text in pages):
            return pages
        return _ocr_pdf(path)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        return [(None, _ocr_image(path))]
    if suffix in {".md", ".txt"}:
        return [(None, path.read_text(encoding="utf-8", errors="ignore"))]
    if suffix in {".docx", ".doc"}:
        from docx import Document as WordDocument

        doc = WordDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        return [(None, "\n".join(paragraphs))]
    if suffix in {".xlsx", ".xls", ".csv"}:
        import pandas as pd

        chunks = []
        if suffix == ".csv":
            df = pd.read_csv(path)
            chunks.append((None, _frame_to_text("CSV", df)))
        else:
            sheets = pd.read_excel(path, sheet_name=None)
            for sheet_name, df in sheets.items():
                chunks.append((None, _frame_to_text(sheet_name, df)))
        return chunks
    raise ValueError(f"暂不支持的文件类型: {suffix}")


def _ocr_image(path: Path) -> str:
    try:
        from PIL import Image
    except ImportError as exc:
        raise ValueError("OCR 依赖未安装，请安装 Pillow") from exc
    return _ocr_image_from_pil(Image.open(path))


def _ocr_pdf(path: Path) -> list[tuple[Optional[int], str]]:
    try:
        from pdf2image import convert_from_path
    except ImportError as exc:
        raise ValueError("扫描 PDF 需要 OCR 支持，请安装 pdf2image、Pillow、pytesseract，并配置 Poppler/Tesseract") from exc
    pages = []
    try:
        poppler_dir = Path(r"C:\Users\Y\.cache\codex-runtimes\codex-primary-runtime\dependencies\bin\override")
        kwargs = {"poppler_path": str(poppler_dir)} if (poppler_dir / "pdftoppm.cmd").exists() else {}
        for index, image in enumerate(convert_from_path(str(path), dpi=220, **kwargs), 1):
            pages.append((index, _ocr_image_from_pil(image)))
    except Exception as exc:
        raise ValueError(f"扫描 PDF OCR 失败，请确认 Poppler/Tesseract 已安装: {exc}") from exc
    return pages


def _ocr_image_from_pil(image) -> str:
    try:
        import pytesseract
    except ImportError as exc:
        raise ValueError("OCR 依赖未安装，请安装 pytesseract") from exc
    tesseract_path = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    if tesseract_path.exists():
        pytesseract.pytesseract.tesseract_cmd = str(tesseract_path)
    tessdata_dir = Path("data/tessdata")
    if tessdata_dir.exists():
        os.environ["TESSDATA_PREFIX"] = str(tessdata_dir.resolve())
    try:
        return pytesseract.image_to_string(image, lang="chi_sim+eng").strip()
    except Exception as exc:
        raise ValueError(f"OCR 识别失败: {exc}") from exc


def load_web_page(url: str) -> tuple[str, str]:
    safe_url = validate_web_url(url)
    response = requests.get(
        safe_url,
        timeout=20,
        headers={"User-Agent": "Mozilla/5.0 (compatible; RAG-KB/1.0)"},
        stream=True,
    )
    response.raise_for_status()
    chunks: list[bytes] = []
    total = 0
    for chunk in response.iter_content(chunk_size=65536):
        if not chunk:
            continue
        total += len(chunk)
        if total > settings.max_web_page_bytes:
            raise ValueError("Web page is too large to import")
        chunks.append(chunk)
    content = b"".join(chunks)
    encoding = response.encoding or response.apparent_encoding or "utf-8"
    html = content.decode(encoding, errors="ignore")
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else safe_url
    blocks = [title, f"URL: {safe_url}"]
    for tag in soup.find_all(["h1", "h2", "h3", "p", "li", "td", "th"]):
        text = tag.get_text(" ", strip=True)
        if text:
            blocks.append(text)
    return title, "\n".join(blocks)


def _frame_to_text(name: str, df) -> str:
    df = df.fillna("")
    lines = [f"Sheet: {name}", "Columns: " + ", ".join(str(col) for col in df.columns)]
    for index, row in df.iterrows():
        cells = [f"{col}: {row[col]}" for col in df.columns if str(row[col]).strip()]
        if cells:
            lines.append(f"Row {index + 1}: " + " | ".join(cells))
    return "\n".join(lines)
