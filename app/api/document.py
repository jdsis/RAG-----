from __future__ import annotations

import json
import mimetypes
import re
from datetime import datetime
from pathlib import Path
from typing import List, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import get_db
from app.deps import get_current_user, require_permission
from app.models import Category, Document, DocumentChunk, DocumentVersion, KnowledgeBase, ParseTask, QaLog, User
from app.schemas import ChangeDocumentCategory, DocumentOut, WebDocumentCreate
from app.services.document_service import (
    assign_document_category,
    create_web_document,
    delete_document_data,
    reclassify_document,
)
from app.services.upload_utils import safe_unique_filename, save_upload_file
from app.services.ops_service import log_audit
from app.services.task_worker import enqueue_parse_task

router = APIRouter(prefix="/documents", tags=["documents"])

SUPPORTED_SUFFIXES = {".pdf", ".txt", ".md", ".docx", ".doc", ".xlsx", ".xls", ".csv", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
PREVIEW_CHUNK_LIMIT = 80


def _get_kb(db: Session, current_user: User, knowledge_base_id: int) -> KnowledgeBase:
    kb = db.query(KnowledgeBase).filter_by(id=knowledge_base_id, user_id=current_user.id, deleted_at=None).first()
    if not kb:
        raise HTTPException(status_code=404, detail="Knowledge base not found")
    return kb


def _get_document(db: Session, current_user: User, document_id: int) -> Document:
    document = db.query(Document).filter_by(id=document_id, user_id=current_user.id, deleted_at=None).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    return document


def _resolve_document_path(document: Document) -> Path:
    stored = Path(document.file_path)
    candidates = [
        stored,
        Path.cwd() / stored,
        settings.upload_path / f"user_{document.user_id}" / f"kb_{document.knowledge_base_id}" / document.filename,
    ]
    for candidate in candidates:
        try:
            resolved = candidate.resolve()
        except OSError:
            resolved = candidate
        if resolved.exists() and resolved.is_file():
            return resolved
    raise HTTPException(status_code=404, detail="Document file not found")


def _outline_heading_level(text: str) -> int:
    stripped = text.strip()
    if not stripped:
        return 0
    if re.match(r"^第[一二三四五六七八九十百千万0-9]+章\s*\S+", stripped):
        return 1
    if re.match(r"^第[一二三四五六七八九十百千万0-9]+节\s*\S+", stripped):
        return 2
    numbered = re.match(r"^(\d+(?:\.\d+){1,2})\s+\S+", stripped)
    if numbered:
        return min(numbered.group(1).count(".") + 1, 3)
    return 0


def _visual_heading_level(text: str, style_name: str = "") -> int:
    outline_level = _outline_heading_level(text)
    if outline_level:
        return outline_level
    style = style_name.lower()
    if "heading 1" in style or "标题 1" in style:
        return 1
    if "heading 2" in style or "标题 2" in style:
        return 2
    if "heading 3" in style or "标题 3" in style:
        return 3
    if "title" in style or "标题" == style.strip():
        return 1
    return 0


def _text_blocks_from_plain_text(text: str) -> tuple[list[dict], list[str]]:
    blocks: list[dict] = []
    outline: list[dict] = []
    heading_index = 0
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            blocks.append({"type": "space"})
            continue
        outline_level = _outline_heading_level(line)
        if outline_level:
            heading_index += 1
            block_id = f"heading-{heading_index}"
            outline.append({"text": line.strip(), "level": outline_level, "block_id": block_id})
            blocks.append({"type": "heading", "level": outline_level, "outline_level": outline_level, "text": line.strip(), "id": block_id})
        elif "|" in line and line.count("|") >= 2:
            cells = [cell.strip() for cell in line.split("|")]
            blocks.append({"type": "table_row", "cells": cells})
        else:
            blocks.append({"type": "paragraph", "text": line})
    return _merge_table_rows(blocks), outline


def _merge_table_rows(blocks: list[dict]) -> list[dict]:
    merged: list[dict] = []
    table_rows: list[list[str]] = []
    for block in blocks:
        if block["type"] == "table_row":
            table_rows.append(block["cells"])
            continue
        if table_rows:
            merged.append({"type": "table", "rows": table_rows})
            table_rows = []
        merged.append(block)
    if table_rows:
        merged.append({"type": "table", "rows": table_rows})
    return merged


def _preview_from_docx(path: Path) -> tuple[list[dict], list[dict], str]:
    from docx import Document as WordDocument

    doc = WordDocument(str(path))
    blocks: list[dict] = []
    outline: list[dict] = []
    content_parts: list[str] = []
    table_by_element = {table._element: table for table in doc.tables}
    heading_index = 0

    for child in doc.element.body:
        if child.tag.endswith("p"):
            paragraph = next((item for item in doc.paragraphs if item._element is child), None)
            if not paragraph:
                continue
            text = paragraph.text.strip()
            if not text:
                blocks.append({"type": "space"})
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            outline_level = _outline_heading_level(text)
            visual_level = _visual_heading_level(text, style_name)
            if visual_level:
                heading_index += 1
                block_id = f"heading-{heading_index}"
                if outline_level:
                    outline.append({"text": text, "level": outline_level, "block_id": block_id})
                blocks.append({
                    "type": "heading",
                    "level": visual_level,
                    "outline_level": outline_level,
                    "text": text,
                    "id": block_id,
                })
            else:
                blocks.append({"type": "paragraph", "text": text})
            content_parts.append(text)
        elif child.tag.endswith("tbl") and child in table_by_element:
            rows = []
            for row in table_by_element[child].rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(cells)
                    content_parts.append(" | ".join(cells))
            if rows:
                blocks.append({"type": "table", "rows": rows})

    return blocks, outline, "\n".join(content_parts)


def _preview_from_file(document: Document, path: Path, chunks: list[DocumentChunk]) -> tuple[list[dict], list[str], str]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _preview_from_docx(path)

    text = "\n\n".join(chunk.content for chunk in chunks)
    if not text and suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
    blocks, outline = _text_blocks_from_plain_text(text)
    return blocks, outline, text


@router.post("/upload", response_model=DocumentOut)
def upload_document(
    knowledge_base_id: int = Form(...),
    auto_classify: bool = Form(True),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("document:upload")),
):
    _get_kb(db, current_user, knowledge_base_id)
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {suffix}")

    upload_dir = settings.upload_path / f"user_{current_user.id}" / f"kb_{knowledge_base_id}"
    target = upload_dir / safe_unique_filename(file.filename, "document")
    file_size = save_upload_file(file, target, settings.max_upload_bytes)

    document = Document(
        user_id=current_user.id,
        knowledge_base_id=knowledge_base_id,
        filename=target.name,
        file_type=suffix.lstrip("."),
        file_path=str(target),
        file_size=file_size,
        status="pending",
    )
    db.add(document)
    db.flush()
    db.add(
        DocumentVersion(
            user_id=current_user.id,
            document_id=document.id,
            version=1,
            filename=target.name,
            file_path=str(target),
            file_size=file_size,
            change_note="初始上传",
        )
    )
    enqueue_parse_task(db, document, "parse", "文档已上传，等待后台解析", auto_classify=auto_classify)
    log_audit(db, current_user, "document.upload", "document", document.id, {"filename": document.filename})
    db.commit()
    db.refresh(document)
    return document


@router.post("/web", response_model=DocumentOut)
def import_web_document(
    payload: WebDocumentCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("document:upload")),
):
    _get_kb(db, current_user, payload.knowledge_base_id)
    upload_dir = settings.upload_path / f"user_{current_user.id}" / f"kb_{payload.knowledge_base_id}"
    upload_dir.mkdir(parents=True, exist_ok=True)
    document = create_web_document(db, current_user, payload.knowledge_base_id, payload.url, upload_dir)
    db.add(document)
    db.flush()
    enqueue_parse_task(db, document, "web", "网页资料已保存，等待后台解析", auto_classify=payload.auto_classify)
    log_audit(db, current_user, "document.web_import", "document", document.id, {"url": payload.url})
    db.commit()
    db.refresh(document)
    return document


@router.get("", response_model=List[DocumentOut])
def list_documents(
    knowledge_base_id: Optional[int] = None,
    category_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("document:read")),
):
    query = db.query(Document).filter_by(user_id=current_user.id, deleted_at=None)
    if knowledge_base_id:
        query = query.filter_by(knowledge_base_id=knowledge_base_id)
    if category_id:
        query = query.filter_by(category_id=category_id)
    return query.order_by(Document.id.desc()).all()


@router.get("/{document_id}", response_model=DocumentOut)
def get_document(document_id: int, db: Session = Depends(get_db), current_user: User = Depends(require_permission("document:read"))):
    return _get_document(db, current_user, document_id)


@router.get("/{document_id}/preview")
def preview_document(document_id: int, db: Session = Depends(get_db), current_user: User = Depends(require_permission("document:read"))):
    document = _get_document(db, current_user, document_id)
    chunks = (
        db.query(DocumentChunk)
        .filter_by(document_id=document.id, user_id=current_user.id)
        .order_by(DocumentChunk.chunk_index)
        .limit(PREVIEW_CHUNK_LIMIT)
        .all()
    )
    try:
        path = _resolve_document_path(document)
        blocks, outline, text = _preview_from_file(document, path, chunks)
    except HTTPException:
        text = "\n\n".join(chunk.content for chunk in chunks)
        blocks, outline = _text_blocks_from_plain_text(text)
    return {
        "id": document.id,
        "filename": document.filename,
        "file_type": document.file_type,
        "knowledge_base_id": document.knowledge_base_id,
        "category_id": document.category_id,
        "category_name": document.category_name,
        "status": document.status,
        "file_size": document.file_size,
        "chunk_count": document.chunk_count,
        "classify_status": document.classify_status,
        "classify_confidence": document.classify_confidence,
        "classify_reason": document.classify_reason,
        "created_at": document.created_at,
        "content": text[:30000],
        "outline": outline[:40],
        "blocks": blocks[:400],
    }


@router.get("/{document_id}/download")
def download_document(document_id: int, db: Session = Depends(get_db), current_user: User = Depends(require_permission("document:download"))):
    document = _get_document(db, current_user, document_id)
    path = _resolve_document_path(document)
    media_type = mimetypes.guess_type(document.filename)[0] or "application/octet-stream"
    disposition = f"attachment; filename*=UTF-8''{quote(document.filename)}"
    return FileResponse(path, media_type=media_type, headers={"Content-Disposition": disposition})


@router.get("/{document_id}/related-qa")
def related_qa(document_id: int, db: Session = Depends(get_db), current_user: User = Depends(require_permission("document:read"))):
    document = _get_document(db, current_user, document_id)
    logs = (
        db.query(QaLog)
        .filter(
            QaLog.user_id == current_user.id,
            QaLog.knowledge_base_id == document.knowledge_base_id,
        )
        .order_by(QaLog.id.desc())
        .limit(80)
        .all()
    )

    related = []
    fallback = []
    for log in logs:
        item = {
            "id": log.id,
            "conversation_id": log.conversation_id,
            "question": log.question,
            "answer": log.answer,
            "response_time_ms": log.response_time_ms,
            "is_answered": log.is_answered,
            "created_at": log.created_at,
            "matched": False,
        }
        try:
            sources = json.loads(log.retrieved_chunks or "[]")
        except json.JSONDecodeError:
            sources = []
        item["sources"] = sources
        if any(source.get("document_name") == document.filename for source in sources if isinstance(source, dict)):
            item["matched"] = True
            related.append(item)
        else:
            fallback.append(item)

    return {
        "document_id": document.id,
        "items": (related or fallback[:10])[:20],
        "matched_count": len(related),
    }


@router.put("/{document_id}/category", response_model=DocumentOut)
def change_document_category(
    document_id: int,
    payload: ChangeDocumentCategory,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("document:update")),
):
    document = _get_document(db, current_user, document_id)
    category = None
    if payload.category_id:
        category = db.query(Category).filter_by(
            id=payload.category_id,
            user_id=current_user.id,
            knowledge_base_id=document.knowledge_base_id,
        ).first()
        if not category:
            raise HTTPException(status_code=404, detail="Category not found")
    assign_document_category(db, document, category, classify_status="manual")
    db.commit()
    db.refresh(document)
    return document


@router.post("/{document_id}/reparse", response_model=DocumentOut)
def reparse_document(
    document_id: int,
    auto_classify: bool = True,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("document:reparse")),
):
    document = _get_document(db, current_user, document_id)
    document.status = "pending"
    enqueue_parse_task(db, document, "reparse", "文档已加入重建索引队列", auto_classify=auto_classify)
    log_audit(db, current_user, "document.reparse", "document", document.id)
    db.commit()
    db.refresh(document)
    return document


@router.post("/{document_id}/reclassify", response_model=DocumentOut)
def auto_reclassify_document(
    document_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("document:update")),
):
    document = _get_document(db, current_user, document_id)
    try:
        reclassify_document(db, document)
    except Exception as exc:
        document.status = "failed"
        document.error_message = str(exc)
    db.commit()
    db.refresh(document)
    return document


@router.delete("/{document_id}")
def delete_document(document_id: int, db: Session = Depends(get_db), current_user: User = Depends(require_permission("document:delete"))):
    document = _get_document(db, current_user, document_id)
    document.deleted_at = datetime.utcnow()
    document.status = "deleted"
    log_audit(db, current_user, "document.soft_delete", "document", document.id, {"filename": document.filename})
    db.commit()
    return {"ok": True}
